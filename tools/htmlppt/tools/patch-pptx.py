#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""patch-pptx.py —— 在原始 PPTX 上打补丁（只改文字/图片，几何完全不动）

思路：放弃"从 HTML 重新生成 PPTX"（会重排），改为：
  HTML 编辑器导出 changes.json → 按 data-pptx-id 定位原对象 → 原地改内容
几何（x/y/w/h、层级、组合）由原文件保证 100% 不变。

用法：
  python tools/patch-pptx.py 原稿.pptx changes.json 输出.pptx [--report r.json]

changes.json 结构：
{
  "pptx": "原稿.pptx",
  "slides": [
    {"slide": 1, "id": 11, "kind": "text", "text": "新的标题"},
    {"slide": 6, "id": 31, "kind": "image", "dataUrl": "data:image/png;base64,..."},
    {"slide": 3, "kind": "background", "index": 1, "dataUrl": "..."}
  ]
}
"""
from __future__ import annotations

import argparse
import base64
import copy
import io
import json
import sys
from pathlib import Path

from pptx import Presentation

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # group
            yield from iter_shapes(sh.shapes)


def find_shape(slide, sid):
    for sh in iter_shapes(slide.shapes):
        try:
            if sh.shape_id == sid:
                return sh
        except Exception:
            continue
    return None


def set_text_keep_style(shape, text: str):
    """替换文本框内容，但保留第一个 run 的字体/字号/颜色。"""
    if not shape.has_text_frame:
        return False
    tf = shape.text_frame
    # 用第一段的第一个 run 作为样式模板
    tmpl_rpr = None
    tmpl_ppr = None
    tmpl_end = None
    for p in tf.paragraphs:
        if tmpl_ppr is None and p._p.find(f"{A}pPr") is not None:
            tmpl_ppr = copy.deepcopy(p._p.find(f"{A}pPr"))
        for r in p.runs:
            rpr = r._r.find(f"{A}rPr")
            if tmpl_rpr is None and rpr is not None:
                tmpl_rpr = copy.deepcopy(rpr)
        if tmpl_end is None:
            ep = p._p.find(f"{A}endParaRPr")
            if ep is not None:
                tmpl_end = copy.deepcopy(ep)
    lines = str(text).split("\n")
    # 清空所有段落，只留一个
    txBody = tf._txBody
    for p in list(tf.paragraphs)[1:]:
        txBody.remove(p._p)
    p0 = tf.paragraphs[0]
    for r in list(p0.runs):
        p0._p.remove(r._r)
    # 移除 break
    for br in p0._p.findall(f"{A}br"):
        p0._p.remove(br)
    if tmpl_ppr is not None:
        old = p0._p.find(f"{A}pPr")
        if old is not None:
            p0._p.remove(old)
        p0._p.insert(0, tmpl_ppr)
    for i, line in enumerate(lines):
        p = p0 if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        if tmpl_rpr is not None:
            r = run._r
            old = r.find(f"{A}rPr")
            if old is not None:
                r.remove(old)
            r.insert(0, copy.deepcopy(tmpl_rpr))
        if tmpl_end is not None and i > 0:
            p._p.append(copy.deepcopy(tmpl_end))
    return True


def decode_data_url(data_url: str):
    if not data_url.startswith("data:"):
        return None
    head, b64 = data_url.split(",", 1)
    return base64.b64decode(b64)


def replace_image(shape, data_url: str):
    """替换图片对象的媒体字节（保持对象几何与裁切不变）。"""
    blob = decode_data_url(data_url)
    if blob is None:
        return False, "无效 dataUrl"
    try:
        blipFill = shape._element.find(f"{P}blipFill")
        blip = blipFill.find(f"{A}blip")
        rid = blip.get(R + "embed")
        part = shape.part.related_part(rid)
    except Exception as ex:
        return False, f"找不到图片关系：{ex}"
    pname = str(part.partname).lower()
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(blob))
        buf = io.BytesIO()
        if pname.endswith(".jpg") or pname.endswith(".jpeg"):
            im.convert("RGB").save(buf, format="JPEG", quality=92)
        elif pname.endswith(".gif"):
            im.save(buf, format="GIF")
        else:
            im.save(buf, format="PNG")
        part._blob = buf.getvalue()
    except Exception:
        part._blob = blob
    return True, pname


def replace_background(slide, index: int, data_url: str):
    """替换第 index 张满页图（背景图层）。"""
    blob = decode_data_url(data_url)
    if blob is None:
        return False, "无效 dataUrl"
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    if index - 1 >= len(pics):
        return False, f"该页只有 {len(pics)} 张图片"
    return replace_image(pics[index - 1], data_url)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("changes")
    ap.add_argument("out")
    ap.add_argument("--report")
    args = ap.parse_args()

    prs = Presentation(args.pptx)
    changes = json.loads(Path(args.changes).read_text(encoding="utf-8"))
    items = changes.get("slides", [])

    report = {"applied": [], "failed": []}
    for c in items:
        try:
            slide = prs.slides[int(c["slide"]) - 1]
        except Exception:
            report["failed"].append({"change": c, "reason": "页不存在"})
            continue
        kind = c.get("kind")
        if kind == "text":
            sh = find_shape(slide, c.get("id"))
            if sh is None:
                report["failed"].append({"change": c, "reason": "对象不存在"})
                continue
            ok = set_text_keep_style(sh, c.get("text", ""))
            report["applied"].append({"slide": c["slide"], "id": c["id"], "kind": "text", "ok": ok})
        elif kind == "image":
            sh = find_shape(slide, c.get("id"))
            if sh is None:
                report["failed"].append({"change": c, "reason": "对象不存在"})
                continue
            ok, info = replace_image(sh, c.get("dataUrl", ""))
            report["applied"].append({"slide": c["slide"], "id": c["id"], "kind": "image", "ok": ok, "part": info})
        elif kind == "background":
            ok, info = replace_background(slide, int(c.get("index", 1)), c.get("dataUrl", ""))
            report["applied"].append({"slide": c["slide"], "kind": "background", "ok": ok, "info": info})
        else:
            report["failed"].append({"change": c, "reason": f"未知类型 {kind}"})

    prs.save(args.out)
    print(f"补丁 {len(report['applied'])} 处 · 失败 {len(report['failed'])} 处 -> {args.out}")
    if args.report:
        Path(args.report).write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    if report["failed"]:
        for f in report["failed"][:5]:
            print("  ✘", f)
        sys.exit(1)


if __name__ == "__main__":
    main()
