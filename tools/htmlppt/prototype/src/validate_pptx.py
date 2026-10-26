#!/usr/bin/env python3
"""validate_pptx.py — independent "can it open" check with python-pptx.

Parses the generated deck with a second, unrelated implementation (python-pptx, not the
library that wrote it) and dumps slide size, shape inventory and text char counts.
"""
import json
import os
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parent.parent
DECK = ROOT / "out" / "deck.pptx"
OUT = (Path(os.environ["HTMLPPT_OUT"]).resolve() if os.environ.get("HTMLPPT_OUT") else ROOT / "out") / "pptx_validate.json"

EXPECTED = {"cx": 12192000, "cy": 6858000, "note": "1280x720 px @96dpi"}


def main():
    report = {"file": str(DECK.relative_to(ROOT)), "ok": False}
    try:
        prs = Presentation(str(DECK))
        slides = []
        total_chars = 0
        for i, slide in enumerate(prs.slides):
            inv = {"slide": i + 1, "shapes": 0, "textFrames": 0, "pictures": 0, "tables": 0,
                   "groups": 0, "autoshapes": 0, "textChars": 0}
            for shape in slide.shapes:
                inv["shapes"] += 1
                if shape.has_text_frame:
                    inv["textFrames"] += 1
                    inv["textChars"] += sum(len(p.text) for p in shape.text_frame.paragraphs)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    inv["pictures"] += 1
                if getattr(shape, "has_table", False) and shape.has_table:
                    inv["tables"] += 1
                    for row in shape.table.rows:
                        for cell in row.cells:
                            inv["textChars"] += len(cell.text)
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    inv["groups"] += 1
                if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                    inv["autoshapes"] += 1
            total_chars += inv["textChars"]
            slides.append(inv)
        size = {"cx": prs.slide_width, "cy": prs.slide_height}
        report.update(
            {
                "ok": len(slides) > 0 and size["cx"] == EXPECTED["cx"] and size["cy"] == EXPECTED["cy"],
                "engine": "python-pptx %s" % __import__("pptx").__version__,
                "slides": len(slides),
                "slideSizeEMU": size,
                "expected": EXPECTED,
                "slideSizeOk": size["cx"] == EXPECTED["cx"] and size["cy"] == EXPECTED["cy"],
                "totalTextChars": total_chars,
                "slidesDetail": slides,
                "notes": "independent parser: if this loads and inventories shapes, the package is structurally valid",
            }
        )
        print(f"[validate] python-pptx opened deck: {len(slides)} slides, {total_chars} text chars, "
              f"size {size['cx']}x{size['cy']} EMU")
    except Exception as exc:  # noqa: BLE001
        report["error"] = f"{type(exc).__name__}: {exc}"
        print(f"[validate] FAILED: {report['error']}", file=sys.stderr)
        OUT.write_text(json.dumps(report, indent=1, ensure_ascii=False))
        sys.exit(1)
    OUT.write_text(json.dumps(report, indent=1, ensure_ascii=False))


if __name__ == "__main__":
    main()
